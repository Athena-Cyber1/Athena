"""Ingestion de fichiers : détection de type → extraction → chunks → index symboles.

Formats traités : texte/code/markdown/json/csv (natif), zip/tar (listing + entrées texte),
pdf (pypdf si dispo), docx (python-docx si dispo), xlsx (openpyxl si dispo), images (métadonnées).
Symboles : Python via ast, JS/Java/C par regex définitions (première classe pour le RAG).
"""
from __future__ import annotations

import ast as pyast
import hashlib
import json
import re
import tarfile
import uuid
import zipfile
from pathlib import Path
from typing import Any

from ..memory import store as memoire

MAX_OCTETS = 20 * 1024 * 1024
CHUNK = 900
CHEVAUCHEMENT = 120

# v10.6 (F14) : les exécutables ne sont PAS des documents — rejet explicite
# (un « strings » de binaire produisait des faux « lignes » et polluait le RAG).
EXTENSIONS_EXECUTABLES = {
    ".exe", ".dll", ".msi", ".sys", ".bin", ".so", ".dmg", ".app",
    ".dylib", ".com", ".scr", ".cpl", ".ocx", ".efi", ".ko", ".o", ".a",
}


def detecter_type(nom: str, octets: bytes) -> dict[str, str]:
    mime = "application/octet-stream"
    if octets.startswith(b"%PDF"):
        mime = "application/pdf"
    elif octets.startswith(b"PK\x03\x04"):
        mime = "application/zip"
    elif octets.startswith((b"\x1f\x8b", b"ustar")):
        mime = "application/gzip"
    elif octets.startswith((b"\x89PNG", b"\xff\xd8\xff")):
        mime = "image/" + ("png" if octets.startswith(b"\x89PNG") else "jpeg")
    elif octets.startswith(b"\x7fELF"):
        mime = "application/x-elf"
    elif octets.startswith(b"MZ"):
        mime = "application/x-pe"
    else:
        try:
            octets.decode("utf-8")
            mime = "text/plain"
        except UnicodeDecodeError:
            pass
    ext = Path(nom).suffix.lower()
    langue = {"py": "python", "js": "javascript", "ts": "typescript", "java": "java", "c": "c",
              "cpp": "cpp", "h": "c", "md": "markdown", "json": "json", "csv": "csv",
              "txt": "texte", "html": "html"}.get(ext.lstrip("."), "")
    return {"mime": mime, "langue": langue, "ext": ext}


def _extraire_texte(nom: str, octets: bytes, mime: str, ext: str) -> dict[str, Any]:
    if mime == "text/plain" or mime in ("application/json",):
        return {"texte": octets.decode("utf-8", errors="replace"), "notes": []}
    if mime == "application/pdf":
        try:
            import io
            from pypdf import PdfReader
            lecteur = PdfReader(io.BytesIO(octets))
            texte = "\n".join((p.extract_text() or "") for p in lecteur.pages)
            return {"texte": texte, "notes": [f"PDF {len(lecteur.pages)} pages"]}
        except ImportError:
            return {"texte": "", "notes": ["PDF : pypdf non installé, texte non extrait"]}
    if ext == ".docx":
        try:
            import io
            import docx
            d = docx.Document(io.BytesIO(octets))
            texte = "\n".join(p.text for p in d.paragraphs)
            return {"texte": texte, "notes": ["DOCX analysé"]}
        except ImportError:
            return {"texte": "", "notes": ["DOCX : python-docx non installé"]}
    if ext in (".xlsx", ".xlsm"):
        try:
            import io
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(octets), read_only=True, data_only=True)
            lignes = []
            for feuille in wb.worksheets:
                lignes.append(f"## feuille {feuille.title}")
                for ligne in feuille.iter_rows(values_only=True):
                    if any(v is not None for v in ligne):
                        lignes.append(" | ".join("" if v is None else str(v) for v in ligne))
            return {"texte": "\n".join(lignes), "notes": [f"XLSX {len(wb.worksheets)} feuilles"]}
        except ImportError:
            return {"texte": "", "notes": ["XLSX : openpyxl non installé"]}
    if ext == ".csv":
        return {"texte": octets.decode("utf-8", errors="replace"), "notes": ["CSV brut"]}
    if mime == "application/zip":
        notes, morceaux_texte = [], []
        try:
            with zipfile.ZipFile(io := __import__("io").BytesIO(octets)) as zf:
                noms = zf.namelist()
                notes.append(f"ZIP {len(noms)} entrées : {', '.join(noms[:15])}")
                for n in noms[:40]:
                    if not n.endswith((".py", ".js", ".ts", ".java", ".c", ".h", ".md", ".txt", ".json", ".csv")):
                        continue
                    try:
                        contenu = zf.read(n).decode("utf-8", errors="replace")
                        morceaux_texte.append(f"### fichier {n}\n{contenu[:20000]}")
                    except Exception:
                        continue
            return {"texte": "\n".join(morceaux_texte), "notes": notes}
        except Exception as e:
            return {"texte": "", "notes": [f"ZIP illisible : {e}"]}
    if mime == "application/gzip" or ext in (".tar", ".tar.gz", ".tgz"):
        try:
            import io
            notes, morceaux_texte = [], []
            with tarfile.open(fileobj=io.BytesIO(octets), mode="r:*") as tf:
                membres = tf.getmembers()[:60]
                notes.append(f"TAR {len(membres)} entrées")
                for m in membres:
                    if not m.isfile():
                        continue
                    if not m.name.endswith((".py", ".js", ".md", ".txt", ".json", ".c", ".h")):
                        continue
                    f = tf.extractfile(m)
                    if f:
                        morceaux_texte.append(f"### fichier {m.name}\n{f.read().decode('utf-8', 'replace')[:20000]}")
            return {"texte": "\n".join(morceaux_texte), "notes": notes}
        except Exception as e:
            return {"texte": "", "notes": [f"TAR illisible : {e}"]}
    if mime.startswith("image/"):
        return {"texte": "", "notes": [f"Image {mime} ({len(octets)} octets) — OCR non disponible dans cette build"]}
    if mime in ("application/x-elf", "application/x-pe"):
        chaine = re.findall(rb"[ -~]{6,}", octets)
        extraits = b"\n".join(chaine[:400]).decode("ascii", errors="replace")
        return {"texte": extraits, "notes": [f"Binaire {mime} — extraction strings"]}
    return {"texte": "", "notes": ["type non extractible"]}


def _symboles(texte: str, langue: str) -> list[str]:
    syms: list[str] = []
    if langue == "python":
        try:
            arbre = pyast.parse(texte)
            for n in pyast.walk(arbre):
                if isinstance(n, (pyast.FunctionDef, pyast.AsyncFunctionDef, pyast.ClassDef)):
                    syms.append(n.name)
        except SyntaxError:
            pass
    else:
        for motif in (r"\bdef\s+(\w+)", r"\bfunction\s+(\w+)", r"\bclass\s+(\w+)",
                      r"\b(?:void|int|char|float|double|bool|public|private)\s+(\w+)\s*\("):
            syms.extend(re.findall(motif, texte))
    return list(dict.fromkeys(syms))[:400]


def _morceaux(texte: str, langue: str) -> list[dict[str, Any]]:
    morceaux = []
    if not texte:
        return morceaux
    for i, debut in enumerate(range(0, len(texte), CHUNK - CHEVAUCHEMENT)):
        bout = texte[debut: debut + CHUNK]
        morceaux.append({"idx": i, "texte": bout, "symboles": _symboles(bout, langue)})
        if i >= 300:
            break
    return morceaux


def ingester(nom: str, octets: bytes) -> dict[str, Any]:
    if len(octets) > MAX_OCTETS:
        return {"statut": "ERREUR", "raison": "fichier trop volumineux (> 20 Mo)"}
    dt = detecter_type(nom, octets)
    # v10.6 (F14) : exécutables (PE/ELF ou extension binaire connue) → REFUS.
    # Le contenu n'est pas un document : l'indexer comme « lignes de texte »
    # fabriquait des documents fantômes et biaisait le RAG.
    if dt["mime"] in ("application/x-pe", "application/x-elf") or dt["ext"] in EXTENSIONS_EXECUTABLES:
        return {
            "statut": "ERREUR",
            "raison": "binaire exécutable non pris en charge — seuls les documents texte, code, archives, PDF et bureautique sont indexés",
            "mime": dt["mime"],
        }
    ex = _extraire_texte(nom, octets, dt["mime"], dt["ext"])
    syms = _symboles(ex["texte"], dt["langue"])
    file_id = uuid.uuid4().hex[:12]
    rec = {
        "file_id": file_id, "nom": nom, "mime": dt["mime"], "sha256": hashlib.sha256(octets).hexdigest(),
        "langue": dt["langue"], "lignes": (ex["texte"].count("\n") + 1) if ex["texte"] else 0,
        "symboles": syms, "statut": "indexed", "taille": len(octets), "erreur": None,
    }
    morceaux = _morceaux(ex["texte"], dt["langue"])
    memoire.enregistrer_fichier(rec, morceaux)
    return {
        "statut": "SUPPORTED" if (ex["texte"] or syms) else "UNKNOWN",
        "file_id": file_id, "nom": nom, "filename": nom, "mime": dt["mime"], "langue": dt["langue"],
        "lignes": rec["lignes"], "symboles": syms[:30], "chunks": len(morceaux),
        "status": "indexed", "kind": memoire._genre(dt["mime"]), "size_bytes": len(octets),
        "notes": ex["notes"], "preuve": f"indexé en mémoire (sha256 {rec['sha256'][:12]}…)",
    }
